"""Physical-coordinate EP6 adapter of results_v2/fix_ep9_input.py; seed 42."""
import os
for k in ['OPENBLAS_NUM_THREADS','OMP_NUM_THREADS','MKL_NUM_THREADS']: os.environ[k]='2'
import sys,csv,time,zipfile,hashlib,importlib.util,json
from pathlib import Path
from xml.etree import ElementTree as ET
import numpy as np
import pandas as pd
from openpyxl import load_workbook
sys.dont_write_bytecode=True
OUT=Path(__file__).resolve().parent; ROOT=OUT.parent; SEED=42
RAW=ROOT.parent/'VoteWatch-EP-voting-data_2004-2022/EP6_RCVs_2022_06_13.xlsx'
EXPORT=ROOT.parent/'export.zip'
def save(name,rows,columns=None):
    with (OUT/name).open('x',newline='',encoding='utf-8') as f:pd.DataFrame(rows,columns=columns).to_csv(f,index=False)
def col(ref):
    n=0
    for c in ref:
        if not c.isalpha():break
        n=n*26+ord(c)-64
    return n-1
def main():
    start=time.perf_counter(); np.random.seed(SEED)
    if not RAW.exists():raise FileNotFoundError(f'REQUIRED RAW FILE: {RAW}')
    assert not (OUT/'input_EP6_fixed.npz').exists(),'No overwriting fixed input'
    spec=importlib.util.spec_from_file_location('legacy_readonly',ROOT/'ep_topological_pipeline.py');legacy=importlib.util.module_from_spec(spec);spec.loader.exec_module(legacy)
    prior=np.load(ROOT/'results/input_EP6.npz');members=prior['members']; index={int(v):i for i,v in enumerate(members)}
    with zipfile.ZipFile(RAW) as z:
        strings=legacy.shared_strings(z)
        for _,e in ET.iterparse(z.open('xl/worksheets/sheet1.xml'),events=('end',)):
            if e.tag!=legacy.NS+'row':continue
            if int(e.attrib['r'])==2:
                pairs=[(col(c.attrib['r']),int(float(legacy.value(c,strings)))) for c in e if col(c.attrib['r'])>=10 and legacy.value(c,strings) not in (None,'')];break
            e.clear()
    pairs.sort();columns=np.array([c for c,v in pairs]);vote_ids=np.array([v for c,v in pairs]);assert len(np.unique(vote_ids))==len(vote_ids)
    states=np.zeros((len(members),len(vote_ids)),np.int8);stored=np.zeros_like(states,bool)
    groups=np.full(len(members),'UNKNOWN',dtype='<U40');rawgroups=np.full(len(members),'',dtype='<U100');sources=np.full(len(members),'physical_H',dtype='<U100')
    names=np.full(len(members),'',dtype='<U120');countries=np.full(len(members),'',dtype='<U80')
    rows=[];missing=[];limitations=[];code_counts={};seen=set()
    with zipfile.ZipFile(RAW) as z:
        strings=legacy.shared_strings(z)
        for _,e in ET.iterparse(z.open('xl/worksheets/sheet1.xml'),events=('end',)):
            if e.tag!=legacy.NS+'row':continue
            rn=int(e.attrib['r']); row={col(c.attrib['r']):legacy.value(c,strings) for c in e.findall(legacy.NS+'c')};e.clear()
            if rn<=3 or row.get(0) in (None,''):continue
            mid=int(float(row[0]));assert mid in index and mid not in seen;seen.add(mid);i=index[mid]
            raw=str(row.get(7) or '').strip();rawgroups[i]=raw;names[i]=str(row.get(3) or '');countries[i]=str(row.get(5) or '')
            if raw:
                groups[i]=legacy.group_name(raw)
                if groups[i]=='NI' and raw.upper() not in ['NI','NA','NON-ATTACHED','NON ATTACHED','NON-ATTACHED MEMBERS']:
                    groups[i]='UNKNOWN';sources[i]='unrecognized_physical_H';limitations.append(dict(item='unrecognized_group',member_id=mid,status='unresolved',detail=raw))
            else:sources[i]='missing_physical_H'
            absent=sorted(set(range(10))-set(row))
            if absent:missing.append(dict(member_id=mid,worksheet_row=rn,missing_columns_zero_based=','.join(map(str,absent)),raw_group=raw))
            for j,c in enumerate(columns):
                value=row.get(int(c))
                if value in (None,''):continue
                code=int(float(value));stored[i,j]=True;code_counts[code]=code_counts.get(code,0)+1
                if code in (1,2,3):states[i,j]=code
                elif code not in (0,4,5,6):limitations.append(dict(item='unknown_vote_code',member_id=mid,status='not_present',detail=f'vote={vote_ids[j]};code={code}'))
            rows.append(dict(member_id=mid,worksheet_row=rn,raw_group=raw))
            if len(rows)%100==0:print(f'[INFO] input EP6 physical rows {len(rows)}/{len(members)}; seed=42',flush=True)
    assert len(seen)==len(members)
    if np.any(groups=='UNKNOWN'):
        memberships=[]
        if EXPORT.exists():
            with zipfile.ZipFile(EXPORT) as z:memberships=list(csv.DictReader(__import__('io').TextIOWrapper(z.open('group_memberships.csv'),encoding='utf-8')))
        for i in np.flatnonzero(groups=='UNKNOWN'):
            labels={legacy.group_name(r['group_code']) for r in memberships if int(r['member_id'])==int(members[i]) and r['term']=='6'}
            if len(labels)==1 and 'NI' not in labels:groups[i]=labels.pop();sources[i]='export.zip/group_memberships.csv;term=6'
            else:limitations.append(dict(item='missing_EP6_group',member_id=int(members[i]),status='explicit_UNKNOWN',detail=f'No unambiguous physical/term-6 label; present_votes={int((states[i]>0).sum())}'))
    yes=states==1;no=states==2;abstain=states==3;present=states>0
    assert not np.any((yes&no)|(yes&abstain)|(no&abstain));assert np.array_equal(present,np.isin(states,[1,2,3]))
    checks=[]
    def check(label,test,count=0):
        checks.append(dict(check=label,passed=bool(test),disagreements=int(count)));assert test,label
    nv=prior['coordinate'].shape[1];check('old vote IDs physical prefix',np.array_equal(vote_ids[:nv],prior['vote_ids']))
    diff=int(np.count_nonzero(states[:,:nv]!=prior['coordinate']));check('shared fixed states versus independent coordinate reconstruction',diff==0,diff)
    rng=np.random.default_rng(SEED);sample=np.unique(np.r_[0,len(vote_ids)-1,rng.integers(0,len(vote_ids),20)])
    workbook=load_workbook(RAW,read_only=True,data_only=True);ws=workbook.worksheets[0];bad=0;tested=0
    for values in ws.iter_rows(min_row=4,max_row=53,values_only=True):
        if values[0] in (None,''):continue
        i=index[int(values[0])]
        for j in sample:
            v=values[int(columns[j])] if int(columns[j])<len(values) else None;code=int(v) if v not in (None,'') else 0;expected=code if code in (1,2,3) else 0
            bad+=int(states[i,j]!=expected);tested+=1
        bad+=int(rawgroups[i]!=str(values[7] or '').strip());tested+=1
    workbook.close();check(f'openpyxl physical coordinates ({tested} comparisons)',bad==0,bad)
    new_yes_dis=int(np.count_nonzero(yes[:,:nv]!=(prior['coordinate']==1)));check('fixed Yes versus physical Yes',new_yes_dis==0,new_yes_dis)
    metrics=[dict(metric='EP6 packed_vs_physical_Yes_disagreements_before',value=int(np.count_nonzero((prior['current']==1)!=(prior['coordinate']==1))),definition='Stored packed alignment versus independent physical coordinates; shared IDs'),dict(metric='EP6 fixed_vs_physical_Yes_disagreements_after',value=new_yes_dis,definition='Fixed Yes versus independent physical coordinates'),dict(metric='EP6_roll_calls',value=len(vote_ids),definition='Physical vote-ID columns in raw EP6 workbook'),dict(metric='EP6_stored_vote_cells',value=int(stored.sum()),definition='Nonempty physical raw cell; NOT presence'),dict(metric='EP6_present_cells',value=int(present.sum()),definition='Raw Yes/No/Abstain only, codes 1/2/3'),dict(metric='EP6_stored_but_not_present',value=int(np.sum(stored&~present)),definition='Nonempty absent/not-voting code; excluded from co-presence'),dict(metric='EP6_present_without_stored_cell',value=int(np.sum(present&~stored)),definition='Must be zero'),dict(metric='UNKNOWN_members',value=int(np.sum(groups=='UNKNOWN')),definition='UNKNOWN remains separate from NI')]
    save('input_presence_audit_stage0.csv',metrics);save('input_presence_audit.csv',metrics);save('input_validation.csv',checks)
    save('vote_codes.csv',[dict(raw_code=k,count=v,state='Yes' if k==1 else 'No' if k==2 else 'Abstain' if k==3 else 'not_present') for k,v in sorted(code_counts.items())])
    save('limitations.csv',limitations,['item','member_id','status','detail']);save('missing_metadata_cells.csv',missing,['member_id','worksheet_row','missing_columns_zero_based','raw_group'])
    md=pd.DataFrame(rows).set_index('member_id').loc[members].reset_index();md['fixed_group']=groups;md['old_group']=prior['groups'];md['group_source']=sources;md['present_votes']=present.sum(1);save('member_metadata_audit.csv',md)
    with (OUT/'input_EP6_fixed.npz').open('xb') as f:np.savez_compressed(f,states=states,yes=yes,no=no,abstain=abstain,present=present,stored_vote_cells=stored,groups=groups,raw_groups=rawgroups,group_source=sources,members=members,vote_ids=vote_ids,physical_columns_zero_based=columns,countries=countries,names=names,seed=np.array(SEED),missing_group=groups=='UNKNOWN')
    with (OUT/'input_parameters.json').open('x') as f:json.dump(dict(UNKNOWN_members=int(np.sum(groups=='UNKNOWN')),UNKNOWN_present_votes=int(present[groups=='UNKNOWN'].sum()),stored_but_not_present=int(np.sum(stored&~present)),members=len(members),votes=len(vote_ids)),f,indent=2)
    save('stage_runtime.csv',[dict(stage='0_input',seconds=time.perf_counter()-start,seed=SEED)])
    print(pd.DataFrame(metrics).to_string(index=False),flush=True)
if __name__=='__main__':main()
